"""Convert the laboratory report Markdown to a readable DOCX."""

import pathlib
import re

from docx import Document
from docx.shared import Cm, Pt


ROOT = pathlib.Path(__file__).parent


def clean(text):
    text = re.sub(r"\[([^]]+)\]\([^)]+\)", r"\1", text)
    return text.replace("**", "").replace("`", "")


def main():
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(1.8)
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    for name in ("Title", "Heading 1", "Heading 2"):
        styles[name].font.name = "Arial"

    lines = (ROOT / "report.md").read_text().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("# "):
            document.add_heading(clean(line[2:]), 0)
        elif line.startswith("## "):
            document.add_heading(clean(line[3:]), 1)
        elif line.startswith("### "):
            document.add_heading(clean(line[4:]), 2)
        elif line.startswith("!["):
            match = re.match(r"!\[([^]]+)\]\(([^)]+)\)", line)
            if match:
                path = ROOT / match.group(2)
                if path.exists():
                    document.add_picture(str(path), width=Cm(15.5))
                    document.add_paragraph(match.group(1), style="Caption")
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                parts = [clean(x.strip()) for x in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"[-: ]+", x or "-") for x in parts):
                    rows.append(parts)
                i += 1
            if rows:
                cols = max(len(row) for row in rows)
                table = document.add_table(rows=len(rows), cols=cols)
                table.style = "Table Grid"
                for ri, row in enumerate(rows):
                    for ci, value in enumerate(row):
                        table.cell(ri, ci).text = value
                        if ri == 0:
                            for run in table.cell(ri, ci).paragraphs[0].runs:
                                run.bold = True
                document.add_paragraph()
            continue
        elif line.startswith("- "):
            document.add_paragraph(clean(line[2:]), style="List Bullet")
        else:
            paragraph = document.add_paragraph(clean(line))
            paragraph.paragraph_format.space_after = Pt(6)
        i += 1

    output = ROOT / "report.docx"
    document.save(output)
    print(output)


if __name__ == "__main__":
    main()
